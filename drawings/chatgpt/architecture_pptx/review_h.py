"""Final owner typography and entrance-glass presentation edits."""
import io,json,subprocess
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[2]
BASE='819460d'
def old(name):return subprocess.check_output(['git','show',f'{BASE}:drawings/chatgpt/architecture_pptx/{name}'],cwd=ROOT)
p=Presentation(io.BytesIO(old('BORN_FRAGRANCE_Architecture.pptx')))
m=json.loads(old('geometry_manifest.json'))
frozen=[s._element.xml for s in list(p.slides)[8:]]
dimensions={(i,sh.name):(sh.left,sh.top,sh.width,sh.height) for i,s in enumerate(p.slides) for sh in s.shapes if sh.name.startswith('GEOM:')}
for i,s in enumerate(list(p.slides)[:8]):
    for sh in list(s.shapes):
        if not sh.has_text_frame:continue
        if sh.text=='BORN FRAGRANCE  /  ARCHITECTURE + INTERIOR DESIGN':
            s.shapes._spTree.remove(sh._element)
        elif 19<sh.top/36000<25:
            for q in sh.text_frame.paragraphs:
                for run in q.runs:run.text=run.text.upper()
            m['slide_titles'][i]=sh.text
for shape_id,color,opacity in [(241,'D5E8F0',30000),(242,'B7D8E8',35000)]:
    sh=next(sh for sh in p.slides[3].shapes if sh.shape_id==shape_id)
    sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color)
    rgb=sh._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill/{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
    alpha=OxmlElement('a:alpha');alpha.set('val',str(opacity));rgb.append(alpha)
assert frozen==[s._element.xml for s in list(p.slides)[8:]]
assert dimensions=={(i,sh.name):(sh.left,sh.top,sh.width,sh.height) for i,s in enumerate(p.slides) for sh in s.shapes if sh.name.startswith('GEOM:')}
m['revision']='H';m['review_source_commit']=subprocess.check_output(['git','rev-parse',BASE],cwd=ROOT,text=True).strip()
(OUT/'geometry_manifest.json').write_text(json.dumps(m,indent=2))
p.save(OUT/'BORN_FRAGRANCE_Architecture.pptx')
for entry in list(p.slides._sldIdLst)[8:]:p.part.drop_rel(entry.rId);p.slides._sldIdLst.remove(entry)
p.save(OUT/'BORN_FRAGRANCE_Review_H.pptx')
print('Review H: all eight banners removed; titles uppercase; entrance glass tinted.')
