"""Owner review I: approved elevation height chains and restrained depth cues."""
import io,json,subprocess
from pathlib import Path
import numpy as np
from pptx.enum.text import PP_ALIGN
OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[2]
exec(compile((OUT/'build.py').read_text().split("s=sheet('Retail interior")[0],str(OUT/'build.py'),'exec'))
BASE='cb669753b7202261c2c032d5b61f229ed36686a9'
def old(name):return subprocess.check_output(['git','show',f'{BASE}:drawings/chatgpt/architecture_pptx/{name}'],cwd=ROOT)
prs=Presentation(io.BytesIO(old('BORN_FRAGRANCE_Architecture.pptx')))
m=json.loads(old('geometry_manifest.json'))
frozen=[s._element.xml for s in list(prs.slides)[8:]]
dimensions={(i,sh.name):(sh.left,sh.top,sh.width,sh.height) for i,s in enumerate(prs.slides) for sh in s.shapes if sh.name.startswith('GEOM:')}
records=m['annotations']
def replace_call(page,title,point,y,via=None):
    s=prs.slides[page-1];r=next(r for r in records if r['slide']==page and r['title']==title)
    oldpoint=r['leader_mm'][0]
    for sh in list(s.shapes):
        dot=sh.shape_type==1 and abs((sh.left+sh.width/2)/36000-oldpoint[0])<.02 and abs((sh.top+sh.height/2)/36000-oldpoint[1])<.02 and sh.width/36000<2
        if sh.name=='LEADER:'+title or dot:s.shapes._spTree.remove(sh._element)
        elif sh.name.startswith('ANNOTATION:'+title):sh.top+=Mm(y-r['box_mm'][1])
    x=r['box_mm'][0];pts=[point]+(via or [])+[(x-9,y+2),(x-3,y+2)]
    for a,b in zip(pts,pts[1:]):line(s,*a,*b,GREY,.55).name='LEADER:'+title
    circle(s,*point,.5,GREY,GREY)
    r['leader_mm']=pts;r['box_mm'][1]=y

# Keep the existing descriptions clear of the new external dimension chain.
v=View(prs.slides[3],65,237,20)
replace_call(4,'BLACK STEEL FRAME',v.p(2387.5,2500),111)
replace_call(4,'FOLDED DISPLAY TRAY',v.p(2175,1800),152)
replace_call(4,'CLEAR TOUGHENED GLASS',v.p(1700,800),190)
v=View(prs.slides[4],65,222,20)
replace_call(5,'UNLIT BLACK LOGO',v.p(1200,1970),111)
replace_call(5,'COUNTER FRONT',v.p(1030,750),178)

height_chains=[]
units=SPEC['confirmed']['units'];sf=SPEC['confirmed']['storefront']
levels=[0,units['plinth_h']]+units['shelf_levels']+[units['frame_h']]
for pos,origin,maxheight in [(3,(65,237),sf['overall_h']),(4,(65,222),units['frame_h'])]:
    s=prs.slides[pos];v=View(s,*origin,20)
    heights=levels+([sf['overall_h']] if pos==3 else [])
    # AP-06 previously had just the fascia height; replace it with the full chain.
    if pos==3:
        for sh in list(s.shapes):
            if 276<=sh.shape_id<=281:s.shapes._spTree.remove(sh._element)
    for lo,hi in zip(heights,heights[1:]):
        before=len(s.shapes);v.dim(W,lo,W,hi,11)
        for sh in list(s.shapes)[before:]:
            sh.name='HEIGHT-CHAIN:'+str(hi-lo)
            if sh.has_text_frame:
                sh.left=Mm(198);sh.width=Mm(11)
                sh.text_frame.paragraphs[0].alignment=PP_ALIGN.LEFT
                for q in sh.text_frame.paragraphs:q.space_after=Pt(0)
    height_chains.append({'slide':pos+1,'levels_mm':heights,'total_mm':maxheight})
    assert sum(b-a for a,b in zip(heights,heights[1:]))==maxheight
# Door height is approved independently of the unit pitch chain.
sh=next(sh for sh in prs.slides[3].shapes if sh.name=='ANNOTATION:CLEAR TOUGHENED GLASS description')
for q in sh.text_frame.paragraphs:
    if q.runs:q.runs[0].text='1000 × 2580 leaf · 12 mm glass'

def shadow(sh):
    sp=sh._element.spPr
    for el in list(sp.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')):sp.remove(el)
    eff=OxmlElement('a:effectLst');outer=OxmlElement('a:outerShdw')
    for k,val in {'blurRad':'40000','dist':'18000','dir':'2700000','algn':'tl','rotWithShape':'0'}.items():outer.set(k,val)
    rgb=OxmlElement('a:srgbClr');rgb.set('val','252725');alpha=OxmlElement('a:alpha');alpha.set('val','8000');rgb.append(alpha)
    outer.append(rgb);eff.append(outer);sp.append(eff)
shadow_count=0
for pos,bounds in [(1,(90,82,210,238)),(3,(65,77,185,237)),(4,(65,92,185,222)),(5,(51,92,206.3,222)),(6,(51,92,206.3,222))]:
    x0,y0,x1,y1=bounds
    for sh in prs.slides[pos].shapes:
        x,y,w,h=[n/36000 for n in (sh.left,sh.top,sh.width,sh.height)]
        if sh.shape_type!=1 or not (x0-.1<=x and y0-.1<=y and x+w<=x1+.1 and y+h<=y1+.1):continue
        if sh.fill.type!=1:continue
        if w>10 and (h>6 or .1<h<2.5):shadow(sh);shadow_count+=1

# Native, soft floor-contact shadows in the existing axonometric projection.
# These are presentation graphics, not additional geometry or dimensions.
world={'BLACK STEEL TRAYS':(250,1100,1400),'NEUTRAL BASE DOORS':(251,1100,380),
       'COUNTER':(1430,1702,500),'REAR BRAND WALL':(1200,3105,1900)}
matrix=np.array([list(pt)+[1] for pt in world.values()])
paper=np.array([next(r['leader_mm'][0] for r in records if r['slide']==8 and r['title']==key) for key in world])
transform=np.linalg.solve(matrix,paper)
def project(x,y):return (np.array([x,y,0,1])@transform).tolist()
s=prs.slides[7];pic=next(sh for sh in s.shapes if sh.shape_type==13 and sh.top/36000>50)
for coords in [
    [(250,450),(290,450),(290,2856),(250,2856)],
    [(2110,450),(2150,450),(2150,2856),(2110,2856)],
    [(700,1656),(1760,1656),(1760,1702),(700,1702)],
    [(1700,1702),(1760,1702),(1760,2156),(1700,2156)],
]:
    pts=[project(*a) for a in coords];f=s.shapes.build_freeform(Mm(pts[0][0]),Mm(pts[0][1]))
    f.add_line_segments([(Mm(a),Mm(b)) for a,b in pts[1:]],close=True)
    sh=f.convert_to_shape();sh.name='SHADOW:floor contact';sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(INK);sh.line.fill.background()
    rgb=sh._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill/{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
    alpha=OxmlElement('a:alpha');alpha.set('val','6500');rgb.append(alpha)
    eff=OxmlElement('a:effectLst');soft=OxmlElement('a:softEdge');soft.set('rad','16000');eff.append(soft);sh._element.spPr.append(eff)
    # Keep text and leaders above the cast-shadow graphics.
    pic._element.addnext(sh._element)
assert dimensions=={(i,sh.name):(sh.left,sh.top,sh.width,sh.height) for i,s in enumerate(prs.slides) for sh in s.shapes if sh.name.startswith('GEOM:')}
assert frozen==[s._element.xml for s in list(prs.slides)[8:]]
m['revision']='I';m['review_source_commit']=BASE;m['height_chains']=height_chains;m['depth_shadows']={'native_objects':shadow_count,'axon_floor_contact':4,'opacity_percent':[8,6.5]}
(OUT/'geometry_manifest.json').write_text(json.dumps(m,indent=2))
prs.save(OUT/'BORN_FRAGRANCE_Architecture.pptx')
for entry in list(prs.slides._sldIdLst)[8:]:prs.part.drop_rel(entry.rId);prs.slides._sldIdLst.remove(entry)
prs.save(OUT/'BORN_FRAGRANCE_Review_I.pptx')
print(f'Review I: two full height chains, {shadow_count} subtle object shadows and four axon floor-contact shadows.')
