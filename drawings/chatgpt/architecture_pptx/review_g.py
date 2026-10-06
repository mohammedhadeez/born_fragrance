"""Revision G: direct component descriptions with deliberately routed leaders.

Loads the committed F deck; geometry and the paused archive are preserved.
"""
import io, json, subprocess
from pathlib import Path
from pptx.enum.text import PP_ALIGN
OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[2]
BASE='7ea40202703e10627d1b27f1e76995172ea6140c'
exec(compile((OUT/'build.py').read_text().split("s=sheet('Retail interior")[0],str(OUT/'build.py'),'exec'))
def old(name):return subprocess.check_output(['git','show',f'{BASE}:drawings/chatgpt/architecture_pptx/{name}'],cwd=ROOT)
prs=Presentation(io.BytesIO(old('BORN_FRAGRANCE_Architecture.pptx')))
manifest=json.loads(old('geometry_manifest.json'))
frozen=[s._element.xml for s in list(prs.slides)[8:]]
geometry={ (i,sh.name):sh._element.xml for i,s in enumerate(prs.slides) for sh in s.shapes if sh.name.startswith('GEOM:') }
records=[]
def remove(s,ids):
    for sh in list(s.shapes):
        if sh.shape_id in ids:s.shapes._spTree.remove(sh._element)
def label(s,x,y,w,title,body):
    for yy,h,t,size,bold in [(y,8,title,9,True),(y+8,19,body,8.5,False)]:
        a=txt(s,x,yy,w,h,t,size,INK if bold else GREY,bold=bold)
        a.name='ANNOTATION:'+title+(' title' if bold else ' description')
        for q in a.text_frame.paragraphs:q.space_after=Pt(0);q.line_spacing=1.1
    return [x,y,w,27]
def call(s,point,x,y,w,title,body,side='right',via=None):
    # Horizontal terminal aligns with the first text line, with a 3 mm gap.
    end=(x-3,y+2) if side=='right' else (x+w+3,y+2)
    elbow=(end[0]-6,end[1]) if side=='right' else (end[0]+6,end[1])
    pts=[point]+(via or [])+[elbow,end]
    for a,b in zip(pts,pts[1:]):
        sh=line(s,*a,*b,GREY,.55);sh.name='LEADER:'+title
    circle(s,*point,.5,GREY,GREY)
    box=label(s,x,y,w,title,body)
    records.append({'slide':list(prs.slides).index(s)+1,'title':title,'box_mm':box,'leader_mm':pts})

# Delete only the identified previous annotation groups, including orphan tails.
remove(prs.slides[1],{245,246,247,248,249,250,288})
remove(prs.slides[2],set(range(87,96)))
remove(prs.slides[3],set(range(282,291))|set(range(298,303)))
remove(prs.slides[4],set(range(243,249))|set(range(256,261)))
for i in (5,6):remove(prs.slides[i],set(range(462,468)))

s=prs.slides[1];v=View(s,90,238,20)
call(s,v.p(2375,1900),226,103,61,'SIDE DISPLAY','250 deep shelving\n25 × 25 main steel frame')
call(s,v.p(1700,2100),226,146,61,'COUNTER / STEEL FRAME','1000 × 450 footprint\n20 × 20 black steel')
call(s,v.p(1850,0),226,238,61,'PIVOT GLASS DOOR','1000-wide leaf · 12 glass\nInward opening',via=[(218,241)])

s=prs.slides[2];v=View(s,75,238,20)
call(s,v.p(1750,2100),213,112,66,'BLACK CEILING TRACK','Adjustable spot heads\n3000 K warm white')
call(s,v.p(2180,1450),213,156,66,'CONCEALED SHELF LIGHT','LEDs below shelving\nShown dashed in ceiling plan')
call(s,v.p(1700,225),213,226,66,'ENTRANCE DOWNLIGHTS','Three approved positions\nFinal aiming on site')

s=prs.slides[3];v=View(s,65,237,20)
call(s,v.p(1200,2900),211,79,76,'HALO-LIT BRAND','1000 wide · 25 stand-offs')
call(s,v.p(2387.5,2400),211,117,76,'BLACK STEEL FRAME','25 × 25 × 1.6 main SHS')
call(s,v.p(2175,1800),211,156,76,'FOLDED DISPLAY TRAY','Warm concealed shelf light')
call(s,v.p(1700,800),211,199,76,'CLEAR TOUGHENED GLASS','12 mm · interior visible beyond')

s=prs.slides[4];v=View(s,65,222,20)
call(s,v.p(1200,1970),211,102,68,'UNLIT BLACK LOGO','600 wide · centre at 1900')
call(s,v.p(2125,1400),211,150,68,'BLACK TRAY / WARM LED','250 deep · 25 downstand')
call(s,v.p(1030,750),211,181,68,'COUNTER FRONT','Mesh left / neutral solid right\nPanel split and top TBC #16')

for i in (5,6):
    s=prs.slides[i];v=View(s,51,222,20)
    # Keep the dimension figures outside the pitch line, clear of the leaders.
    for sh in s.shapes:
        if sh.shape_id in {425,431,437,443,449,455,461}:
            sh.left=Mm(216.5);sh.width=Mm(13)
            sh.text_frame.paragraphs[0].alignment=PP_ALIGN.LEFT
    edge=3090 if i==5 else 2630
    call(s,v.p(edge,2500),237,88,43,'EXPANDED MESH','Upper infill\n2200–2600')
    call(s,v.p(edge,1400),237,148,43,'BLACK STEEL TRAY','25 folded downstand\nConcealed warm LED')
    call(s,v.p(edge,25),237,211,43,'RECESSED PLINTH','100 high · 50 recess\nWarm toe-kick light')
    # Replace the cramped text inside the column with a clear nearby description.
    remove(s,{24})
    x=51 if i==5 else 183.8
    a=txt(s,x,70,27,18,'SOLID COLUMN RETURN\nDisplay faces street',8,GREY)
    for q in a.text_frame.paragraphs:q.space_after=Pt(0)
    sh=next(sh for sh in s.shapes if sh.shape_id==485)
    sh.left=Mm(51);sh.top=Mm(276);sh.width=Mm(220)

# Centre horizontal dimension numbers in their boxes; preserve all values.
for s in list(prs.slides)[1:7]:
    for sh in s.shapes:
        if sh.has_text_frame and sh.text.strip().isdigit() and sh.width/36000>=20:
            sh.text_frame.paragraphs[0].alignment=PP_ALIGN.CENTER

# Crop the axonometric's unused raster margin in PowerPoint, without resampling.
s=prs.slides[7]
oldpoints={}
for key,shape_id in [('mesh',10),('tray',15),('base',20),('entry',25),('track',30),('logo',35),('counter',40),('cut',45)]:
    sh=next(a for a in s.shapes if a.shape_id==shape_id)
    oldpoints[key]=((sh.left+sh.width/2)/36000,(sh.top+sh.height/2)/36000)
for sh in list(s.shapes):
    if 49<sh.top/36000<277:s.shapes._spTree.remove(sh._element)
crop=(45,45,1720,2165);left,top,right,bottom=crop
ix,iy,ih=110,57,221;iw=ih*(right-left)/(bottom-top)
pic=s.shapes.add_picture(str(OUT/'axonometric_revE.png'),Mm(ix),Mm(iy),width=Mm(iw),height=Mm(ih))
pic.crop_left=left/2560;pic.crop_right=1-right/2560
pic.crop_top=top/2210;pic.crop_bottom=1-bottom/2210
def anchor(key):
    x,y=oldpoints[key];px=(x-81)*10;py=(y-57)*10
    return ix+(px-left)*iw/(right-left),iy+(py-top)*ih/(bottom-top)
for key,x,y,title,body,side in [
    ('mesh',19,97,'EXPANDED MESH','Upper infill 2200–2600\nApprox. 20 × 40 diamonds','left'),
    ('tray',19,156,'BLACK STEEL TRAYS','250 deep · 3 mm steel\n25 folded front downstand','left'),
    ('base',19,207,'NEUTRAL BASE DOORS','1.5 mm PPC steel\n600 base incl. 100 plinth','left'),
    ('entry',19,249,'STREET DISPLAY CUT','Upper column removed\n1500 clear entrance','left'),
    ('track',315,65,'CEILING TRACKS','Exploded above room\nInstalled at 2600 · 3000 K','right'),
    ('logo',315,121,'REAR BRAND WALL','Light plaster finish\n600-wide unlit native logo','right'),
    ('counter',315,192,'COUNTER','1000 × 450 × 900\nMesh left / neutral panel right','right'),
    ('cut',315,244,'NEAR-SIDE CUTAWAY','Upper joinery omitted at 600\nFull run shown in elevations','right'),
]:call(s,anchor(key),x,y,85 if side=='right' else 75,title,body,side)

assert frozen==[s._element.xml for s in list(prs.slides)[8:]]
assert geometry=={(i,sh.name):sh._element.xml for i,s in enumerate(prs.slides) for sh in s.shapes if sh.name.startswith('GEOM:')}
manifest['revision']='G';manifest['review_source_commit']=BASE
manifest['annotations']=records
(OUT/'geometry_manifest.json').write_text(json.dumps(manifest,indent=2))
prs.save(OUT/'BORN_FRAGRANCE_Architecture.pptx')
for entry in list(prs.slides._sldIdLst)[8:]:prs.part.drop_rel(entry.rId);prs.slides._sldIdLst.remove(entry)
prs.save(OUT/'BORN_FRAGRANCE_Review_G.pptx')
print(f'Revision G: {len(records)} direct annotations; measured geometry and paused archive preserved.')
