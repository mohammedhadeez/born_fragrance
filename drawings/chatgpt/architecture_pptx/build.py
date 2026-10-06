"""Native, editable A3 architectural presentation; millimetres are paper units."""
import json, math, subprocess, hashlib
from pathlib import Path
import cairosvg
from pptx import Presentation
from pptx.util import Mm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[2]
SPEC = json.loads((ROOT/'spec/SHOP_SPEC.json').read_text())
L, C = SPEC['locked'], SPEC['confirmed']
W, D, H = L['internal_size']['W'], L['internal_size']['D'], L['ceiling_height']
SHA = subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()
INK, GREY, GOLD, BEIGE, PAPER, RED, BLUE = '252725','737871','9A825A','E8E1D5','FCFBF7','A44338','466976'
prs = Presentation(); prs.slide_width=Mm(420); prs.slide_height=Mm(297)
checks=[]; index=[]
logo=OUT/'logo.png'
cairosvg.svg2png(url=str(ROOT/L['signage']['logo_asset']),write_to=str(logo),output_width=1400)
from PIL import Image
lw,lh=Image.open(logo).size
ratio=1419.5/643

def txt(s,x,y,w,h,text,size=10,color=INK,font='Liberation Sans',bold=False):
    box=s.shapes.add_textbox(Mm(x),Mm(y),Mm(w),Mm(h)); tf=box.text_frame
    tf.word_wrap=True; tf.margin_left=tf.margin_right=Mm(0);tf.margin_top=tf.margin_bottom=Mm(0)
    for i,t in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=t
        p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(4)
    return box
def rect(s,x,y,w,h,fill=None,color=INK,weight=.6,dash=False):
    a=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Mm(x),Mm(y),Mm(w),Mm(h))
    if fill:a.fill.solid();a.fill.fore_color.rgb=RGBColor.from_string(fill)
    else:a.fill.background()
    a.line.color.rgb=RGBColor.from_string(color);a.line.width=Pt(weight)
    if dash:a.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    return a
def line(s,x,y,u,v,color=INK,weight=.7,dash=False):
    a=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Mm(x),Mm(y),Mm(u),Mm(v))
    a.line.color.rgb=RGBColor.from_string(color);a.line.width=Pt(weight)
    if dash:a.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    return a
def circle(s,x,y,r,color=INK,fill=None):
    a=s.shapes.add_shape(MSO_SHAPE.OVAL,Mm(x-r),Mm(y-r),Mm(2*r),Mm(2*r))
    if fill:a.fill.solid();a.fill.fore_color.rgb=RGBColor.from_string(fill)
    else:a.fill.background()
    a.line.color.rgb=RGBColor.from_string(color);a.line.width=Pt(.7)
def leader(s,x,y,u,v,label,w=75):
    line(s,x,y,u,v,GREY,.45);circle(s,x,y,.65,GREY,GREY)
    txt(s,u+2,v-4,w,12,label,8,GREY)
def cloud(s,x,y,w,h,label):
    # Scalloped closed boundary, kept outside the text field.
    pts=[]
    for a,b in [((x,y),(x+w,y)),((x+w,y),(x+w,y+h)),((x+w,y+h),(x,y+h)),((x,y+h),(x,y))]:
        dx,dy=b[0]-a[0],b[1]-a[1];length=math.hypot(dx,dy);n=max(1,round(length/5))
        for k in range(n):
            for j in range(9):
                t=(k+j/8)/n; bulge=1.2*math.sin(math.pi*j/8)
                pts.append((a[0]+dx*t+dy/length*bulge,a[1]+dy*t-dx/length*bulge))
    f=s.shapes.build_freeform(Mm(pts[0][0]),Mm(pts[0][1]))
    f.add_line_segments([(Mm(a),Mm(b)) for a,b in pts[1:]],close=True)
    sh=f.convert_to_shape();sh.fill.background();sh.line.color.rgb=RGBColor.from_string(RED);sh.line.width=Pt(.6)
    txt(s,x+3,y+3,w-6,h-5,label,8.5,RED)
def notes(s,items):
    y=58
    for title,body in items:
        txt(s,303,y,99,8,title,10,GOLD,bold=True)
        txt(s,303,y+10,98,34,body,10)
        y+=49
def sheet(title,subtitle,scale='NTS',items=None):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(PAPER)
    n=len(prs.slides);code=f'AP-{n:02d}';index.append((code,title,scale))
    txt(s,15,10,295,6,'BORN FRAGRANCE  /  ARCHITECTURE + INTERIOR DESIGN',9,GOLD,bold=True)
    txt(s,15,20,330,14,title,23,font='Liberation Serif')
    txt(s,15,36,345,8,subtitle,9,GREY)
    s.shapes.add_picture(str(logo),Mm(365),Mm(13),width=Mm(38),height=Mm(38/ratio))
    line(s,15,48,405,48,GOLD,.7);line(s,15,272,405,272,GOLD,.7)
    txt(s,15,278,198,10,'FOR COORDINATION  ·  SITE ITEMS OPEN  ·  REV B\n06 OCT 2026  /  SOURCE REV A',7.5,GREY)
    txt(s,214,278,145,10,f'{scale} @ A3  ·  DIMENSIONS mm\nSPEC SNAPSHOT {SHA[:10]}',7.5,GREY)
    txt(s,372,278,32,8,code,12,bold=True)
    if items:
        line(s,293,57,293,260,BEIGE,.6);notes(s,items)
    return s

class View:
    def __init__(self,s,x,y,scale):self.s=s;self.x=x;self.y=y;self.scale=scale
    def p(self,x,y):return self.x+x/self.scale,self.y-y/self.scale
    def r(self,x,y,w,h,fill=None,color=INK,dash=False,name=None):
        u,v=self.p(x,y+h);a=rect(self.s,u,v,w/self.scale,h/self.scale,fill,color,.7,dash)
        if name:
            a.name=f'GEOM:{name}';checks.append({'slide':len(prs.slides),'shape':a.name,'w':w,'h':h,'scale':self.scale})
        return a
    def l(self,x,y,u,v,color=INK,weight=.7,dash=False):
        a,b=self.p(x,y);c,d=self.p(u,v);return line(self.s,a,b,c,d,color,weight,dash)
    def t(self,x,y,t,w=70,size=8,color=INK):
        a,b=self.p(x,y);return txt(self.s,a,b,w,12,t,size,color)
    def dim(self,x,y,u,v,off=8,text=None):
        a,b=self.p(x,y);c,d=self.p(u,v)
        if y==v:
            z=b+off;line(self.s,a,b,a,z,GREY,.4);line(self.s,c,d,c,z,GREY,.4);line(self.s,a,z,c,z,GREY,.5)
            for k in (a,c):line(self.s,k-1,z+1,k+1,z-1,GREY,.5)
            txt(self.s,(a+c)/2-15,z-5,30,5,text or f'{abs(u-x):g}',8,BLUE)
        else:
            z=a+off;line(self.s,a,b,z,b,GREY,.4);line(self.s,c,d,z,d,GREY,.4);line(self.s,z,b,z,d,GREY,.5)
            for k in (b,d):line(self.s,z-1,k+1,z+1,k-1,GREY,.5)
            txt(self.s,z-12,(b+d)/2,24,6,text or f'{abs(v-y):g}',8,BLUE)
    def logo(self,x,y,w):
        a,b=self.p(x,y+w/ratio);self.s.shapes.add_picture(str(logo),Mm(a),Mm(b),width=Mm(w/self.scale),height=Mm(w/ratio/self.scale))
    def bar(self,x,y,length):
        a,b=self.p(x,y);line(self.s,a,b,a+length/self.scale,b,INK,1.5)
        for k in (0,length/2,length):line(self.s,a+k/self.scale,b-1,a+k/self.scale,b+1)
        txt(self.s,a,b+2,length/self.scale+15,6,f'0                         {length:g} mm   |   1:{self.scale}',7,GREY)

def shell(v,plan=True):
    v.r(0,0,W,D if plan else H,name='internal-envelope')
def unitplan(v):
    for x in (0,W-250):
        v.r(x,450,250,D-450,BEIGE)
        for k in range(1,4):v.l(x,450+k*664,x+250,450+k*664)
    for x in (0,1950):v.r(x,0,450,450,BEIGE)
    for x in (250,1650):v.r(x,D-250,500,250,BEIGE)
def plan(v):
    shell(v);unitplan(v);v.r(700,1706,1000,450,'D6D2C8',name='counter-plan')
    v.l(450,0,450+4+488,0,BLUE,1.4);v.l(946,0,1946,0,BLUE,1.4)
    # Inward swing of the long side about right-offset pivot, correct 900 radius.
    px=1846;rr=900
    prev=(px-rr,0)
    for j in range(1,25):
        t=j*math.pi/48;pt=(px-rr*math.cos(t),rr*math.sin(t));v.l(*prev,*pt,GREY,.4,True);prev=pt
    v.l(px,-100,px,900,BLUE,.6,True)

s=sheet('Retail interior / coordinated design','Design intent, measured geometry and assembly strategies',items=[
 ('PROJECT','2400 × 3106 internal\n2600 ceiling\nBlack steel / plaster / stone-look floor'),
 ('ISSUE PURPOSE','Client review and consultant coordination. Detailed strategies are preliminary pending the stated confirmations.'),
 ('DRAWING PRIORITY','Geometry follows approved Rev A. Render conveys material and lighting character only.'),
 ('READING SEQUENCE','Layout → elevations → materials → assembly details → coordination decisions.')])
s.shapes.add_picture(str(ROOT/'assets/renders/born-fragrance-shop.png'),Mm(22),Mm(57),width=Mm(253),height=Mm(221))
# Fit original image to actual aspect ratio within the reserved field, without cropping.
p=s.shapes[-1];im=Image.open(ROOT/'assets/renders/born-fragrance-shop.png');p.height=Mm(200);p.width=Mm(200*im.width/im.height)
txt(s,22,260,260,9,'REFERENCE RENDER · NOT TO SCALE · styling props excluded from technical scope',7,GREY)

s=sheet('Drawing basis / document hierarchy','Confirmed design, assumed site conditions and fabrication decisions',items=[('SOURCE PDF','Repository presentation.pdf, 10 pages. Reorganised and redrawn into an A3 landscape architectural package.'),('SOURCE LIMIT','Named reference books were not supplied. No page-specific technical claims or citations are made.'),('ISSUE STATUS','Final presentation file; preliminary technical details. Not a fabrication or compliance approval.')])
for y,code,title,body,col in [(66,'D','CONFIRMED DESIGN','Black geometry / blue dimensions. Owner-confirmed Rev A, including four 664 modules and the centred counter.',INK),(120,'S','SITE ASSUMPTION','Red clouds #4, #23, #26. No dimensions have been independently surveyed. Internal finish faces and FFL require survey.',RED),(174,'F','FABRICATION PROPOSAL','Red clouds #16 / #26. Detail geometry, tolerances and products require consultant and fabricator confirmation.',RED)]:
    circle(s,27,y+8,7,col);txt(s,24,y+4,10,8,code,11,col,bold=True);txt(s,43,y,240,9,title,12,col,bold=True);txt(s,43,y+12,232,30,body,12)
cloud(s,20,233,260,24,'TBC #16 / #26 — no structural capacities, glass certifications or anchor schedules are inferred.')

s=sheet('Overall plan / setting-out','AP-03 · front-left internal origin; x across, y into shop','1:20',[
 ('01 / ENVELOPE','2400 × 3106 clear finished dimensions assumed. Verify finished faces before fabrication (#4).'),('02 / COUNTER','1000 × 450 × 900H\nx = 700–1700\ny = 1706–2156'),('03 / ACCESS CHECK','450 on either side of counter. These do not meet the project’s assumed 900 route. Consultant decision #16 required.'),('04 / REAR CONNECTION','Side units continue to rear wall. Rear bays occupy x=250–750 and 1650–2150; no footprint overlap.')])
v=View(s,70,240,20);plan(v);v.dim(0,0,W,0,12);v.dim(0,0,0,D,-13);v.dim(250,1706,700,1706,8);v.dim(1700,1706,2150,1706,8)
v.dim(1700,2156,1700,2856,12);v.t(700,1980,'COUNTER',50);v.t(450,3300,'REAR',50);v.bar(0,-480,2000)
cloud(s,204,146,76,27,'#16 — 450 SIDE GAPS\nEscape / staff access unresolved')
txt(s,206,72,70,24,'N ↓ storefront\nASSUMED #23',9,RED)

s=sheet('Floor finish / module setting-out','AP-04 · porcelain extent includes floor beneath units','1:20',[
 ('M04 / PORCELAIN','1200 × 600 rectified stone-look porcelain. Option A: long side across shop. Joint = 2; perimeter = 5.'),('WIDTH CHAIN','5 + 1194 + 2 + 1194 + 5 = 2400'),('DEPTH CHAIN','5 + 343 + (4 × 600) + (5 × 2) + 343 + 5 = 3106'),('FLOOR BUILD-UP','Tile, adhesive and substrate thicknesses TBC #26. Flush threshold detail on AP-17; no step introduced.')])
v=View(s,70,240,20);shell(v);v.r(5,5,2390,3096,None,GREY)
v.l(1200,5,1200,D-5,GREY,.45)
y=348
for i in range(5):v.l(5,y+1,W-5,y+1,GREY,.45);y+=602
v.dim(0,0,W,0,12);v.dim(0,0,0,D,-13);v.dim(5,D-5,1199,D-5,-8);v.dim(0,5,0,348,-6);v.bar(0,-480,2000)
cloud(s,207,186, 70,38,'#4 / #26\nSurvey before tile ordering.\nCuts exclude fabrication tolerances.')

s=sheet('Ceiling / lighting coordination','AP-05 · lighting positions follow confirmed Rev A','1:20',[
 ('L01 / TRACKS','Two black tracks: x=650 / 1750, y=450–2656. Four adjustable cylindrical spots per track. All lighting 3000 K.'),('L02 / DOWNLIGHTS','Three fittings at x=700 / 1200 / 1700; y=225. Spacing retained from approved design.'),('L03 / SHELF + PLINTH','Concealed under-shelf LEDs; toe-kick LEDs 2 × 2656. Access and wiring strategy AP-14.'),('CEILING / MEP','Smooth gypsum, warm off-white. Ceiling support, fire/MEP devices and driver ventilation require consultant coordination #16.')])
v=View(s,70,240,20);shell(v)
for x in C['lighting']['track_x']:
    v.l(x,450,x,2656,INK,2)
    for y in (600,1200,1800,2400):a,b=v.p(x,y);circle(s,a,b,1.8)
for x,y in C['lighting']['downlights']:a,b=v.p(x,y);circle(s,a,b,2,GOLD)
for x in (220,2180):v.l(x,450,x,D,GOLD,1,True)
v.dim(0,0,W,0,12);v.dim(0,0,0,D,-13);v.dim(650,450,650,2656,8);v.bar(0,-480,2000)

s=sheet('Storefront / external elevation','AP-06 · glazed entrance and halo-lit fascia','1:20',[
 ('FRONTAGE','450 + 1500 + 450 = 2400. The 450 × 450 display columns face the street; shelf depth remains 250.'),('G01 / GLAZING','Door 1000 × 2580 × 12 toughened glass. Fixed pane 488 W. Three nominal 4 gaps: 488 + 1000 + 12 = 1500.'),('SIGN BAND','600 high above 2600 ceiling datum; total storefront height 3200. Exterior logo 1000 W, uniformly scaled.'),('DETAIL REFERENCES','Head AP-15\nJamb / pivot AP-16\nFlush threshold AP-17\nSign backing AP-18')])
v=View(s,65,237,20);v.r(0,0,W,3200,name='storefront-envelope');v.r(0,2600,W,600,BEIGE)
for x in (0,1950):
    v.r(x,0,450,2600,None,INK)
    for z in C['units']['shelf_levels']:v.l(x,z,x+450,z)
v.r(454,10,488,2580,None,BLUE);v.r(946,10,1000,2580,None,BLUE,name='door-leaf')
v.logo(700,2600+(600-1000/ratio)/2,1000)
v.dim(0,0,W,0,12);v.dim(0,0,0,3200,-13);v.dim(1946,10,1946,2590,8);v.dim(2400,2600,2400,3200,10)
v.t(1050,1250,'12 mm GLASS',55);v.bar(0,-450,2000)
cloud(s,205,192,74, 30,'#16 — top/bottom clearance\n10 + 2580 + 10 = 2600\nPROPOSED split; glazier to confirm')

s=sheet('Rear focal wall / internal elevation','AP-07 · wall-mounted non-illuminated brand mark','1:20',[
 ('REAR MODULES','250 side + 500 display + 900 logo zone + 500 display + 250 side = 2400. Rear display depth 250.'),('BRAND','600 W logo at centre z=1900. Native asset ratio retained; approximately 271.8 H. No illumination or rear mesh.'),('DISPLAY LEVELS','600 / 1000 / 1400 / 1800 / 2200 AFFL. Frame 2600; top mesh is confined to side units.'),('ASSEMBLY','Removable central back panel and wall fixing strategy AP-19. Rear shelf LEDs return into accessible service routes.')])
v=View(s,65,220,20);shell(v,False)
for x in (0,2150):v.r(x,0,250,2600,BEIGE)
for x in (250,1650):
    v.r(x,0,500,600,BEIGE)
    for z in C['units']['shelf_levels']:v.l(x,z,x+500,z)
    v.l(x,0,x,2600);v.l(x+500,0,x+500,2600)
v.logo(900,1900-300/ratio,600);v.r(700,0,1000,900,'DDD7CC',name='counter-elevation');v.r(1050,150,300,450,None,GREY)
v.dim(0,0,W,0,12);v.dim(0,0,0,H,-13);v.dim(1700,0,1700,900,8);v.bar(0,-500,2000)

def side(title,reverse=False):
    s=sheet(title,'Full side run retained; rear units fit between side footprints','1:20',[
      ('FOUR MODULES','450 storefront + 4 × 664 = 3106. Module boundaries at 450 / 1114 / 1778 / 2442 / 3106.'),('VERTICAL SCHEDULE','100 plinth + 500 cabinet + 5 × 400 = 2600. Display decks at 600 to 2200.'),('MESH / DOORS','Mesh zone 2200–2600 on front plane. Doors 1.5 mm steel, flush handleless, light-neutral PPC.'),('FABRICATION HOLD','664 is a design module, not a shelf cutting size. Shared posts, welds, coatings and adjustment affect clear spans. See AP-12.')])
    v=View(s,47,224,20);v.r(0,0,D,H,name='side-envelope')
    start=0 if reverse else 450;end=2656 if reverse else D
    v.r(2656 if reverse else 0,0,450,H,BEIGE)
    for k in range(5):v.l(start+k*664,0,start+k*664,H)
    v.r(start,0,2656,600,BEIGE);v.r(start,2200,2656,400,None,GREY)
    for z in C['units']['shelf_levels']:v.l(start,z,end,z)
    for k in range(0,2656,100):v.l(start+k,2200,min(start+k+250,end),2600,GREY,.25)
    v.dim(0,0,D,0,15);v.dim(0,0,0,H,-13)
    for k in range(4):v.dim(start+k*664,0,start+(k+1)*664,0,7)
    v.t(0,-450,'REAR' if reverse else 'STOREFRONT');v.t(2500,-450,'STOREFRONT' if reverse else 'REAR');v.bar(0,-750,2000)
side('Left display / internal elevation');side('Right display / internal elevation',True)

s=sheet('Spatial assembly / cutaway axonometric','AP-10 · generated from approved coordinates; parallel projection','NTS',[
 ('SPATIAL LOGIC','Four side modules on each wall; rear bays sit between side-unit faces. Storefront columns turn toward the street.'),('COUNTER','Centred counter retained. Side access conflicts remain visible in plan AP-03 and in the decision register.'),('MODEL LIMIT','Diagrammatic cutaway. Wall thickness, fixings and glass hardware are not inferred from the render.'),('MATERIAL LANGUAGE','Black steel structure; neutral cabinets; plaster backdrop; continuous porcelain floor. No loose products or styling props.')])
def iso(x,y,z):return (139+(x-y)*.031,238-(x+y)*.012-z*.039)
def wirebox(x,y,z,w,d,h,col=INK):
    verts=[(x+i*w,y+j*d,z+k*h) for k in (0,1) for j in (0,1) for i in (0,1)]
    for a,b in [(0,1),(0,2),(1,3),(2,3),(4,5),(4,6),(5,7),(6,7),(0,4),(1,5),(2,6),(3,7)]:line(s,*iso(*verts[a]),*iso(*verts[b]),col,.65)
wirebox(0,0,0,W,D,0,GREY)
for x in (0,2150):
    wirebox(x,450,0,250,2656,600)
    for y in (450,1114,1778,2442,3106):line(s,*iso(x,y,0),*iso(x,y,2600))
    for z in (1000,1400,1800,2200,2600):wirebox(x,450,z,250,2656,0,GREY)
for x in (250,1650):wirebox(x,2856,0,500,250,2600,GREY)
wirebox(700,1706,0,1000,450,900,GOLD)
for x in (0,1950):wirebox(x,0,0,450,450,2600)
txt(s,24,254,260,10,'Axonometric geometry is independent of the AI render. View is not a measurement source.',8,GREY)

s=sheet('Material palette / visible finishes','AP-11 · sample-led specification; colours are indicative','NTS')
materials=[('M01','BLACK STEEL',INK,'25 × 25 × 1.6 main SHS\n20 × 20 × 1.6 secondary SHS\nMatte black powder coat'),('M02','CABINET STEEL','D6D1C6','1.5 mm PPC doors\nLight-neutral finish\nFlush, handleless'),('M03','PLASTER',BEIGE,'Raw light textured plaster\nApprove physical sample\nSubstrate / buildup #26'),('M04','PORCELAIN','C5C0B5','1200 × 600 rectified\n2 grout / 5 perimeter joint\nStone-look, sample approval'),('G01','CLEAR GLASS','DDE8E8','12 mm toughened door\nHardware / gaps #16\nFixed pane specification TBC'),('M05','EXPANDED MESH','66665E','Approx. 20 × 40 diamond\nBlack finish / framed edges\nGauge and fixing #16')]
for i,(code,title,col,body) in enumerate(materials):
    x=20+(i%3)*132;y=61+(i//3)*101
    rect(s,x,y,118, 30,col,col)
    txt(s,x,y+35,120,8,code+' / '+title,12,bold=True);txt(s,x,y+47,118,40,body,10)

s=sheet('Display module / fabrication logic','AP-12 · typical module and installation sequence','1:10',[
 ('MODULE ≠ CUT SIZE','664 is the module allocation. Shared 25 posts are centred at internal boundaries; end posts are inside run ends.'),('UPRIGHT CENTRES','y=462.5 / 1114 / 1778 / 2442 / 3093.5. Clear spans: 626.5 / 639 / 639 / 626.5 before fit allowances.'),('WORKSHOP','Weld frames on jigs; deburr; prepare and coat. Confirm weld size, distortion allowance and coating buildup #16.'),('SITE','Survey → dry fit → level/adjust → anchor to verified floor/wall → connect services → fit removable trays and mesh. No gypsum anchorage.')])
v=View(s,79,250,10);v.r(0,0,664,1800,None,GREY) # lower portion, explicit break
for x in (0,639):v.r(x,0,25,1800,INK)
v.r(0,0,664,600,BEIGE)
for z in (600,1000,1400,1800):v.r(25,z-25,614,25,INK)
v.dim(0,0,664,0,8);v.dim(0,600,0,1800,-10);v.t(-100,1940,'LOWER MODULE / UPPER ZONE CONTINUES',100)
cloud(s,170,93,107,91,'PROPOSED #16\n\nDrawing shows an isolated module.\nDo not copy 614 clear to shared-post bays.\n\nShelf fit clearance and frame segmentation require fabricator approval.\n\nFull elevation: AP-08 / AP-09.')

s=sheet('Shelf / frame / concealed LED','AP-13 · section through 250-deep display shelf','1:2',[
 ('CONFIRMED','3 mm folded MS tray. 25 front downstand. 250 total shelf projection. 20 × 20 × 1.6 support SHS.'),('LIGHTING STRATEGY','Removable aluminium LED channel beneath tray; diffuser screened by downstand. 3000 K. Profile, setback and wattage TBC #16.'),('SERVICE ACCESS','Tray/channel accessible from below. Concealed cable follows rear upright with protected holes, grommets and strain relief.'),('LOAD / FIT','Shelf loading, span deflection, welds, fasteners and actual LED housing dimensions require fabricator/consultant confirmation.')])
v=View(s,64,158,2);v.r(0,0,250,3,INK,name='shelf-section');v.r(247,-22,3,25,INK)
v.r(15,-20,20,20,None,INK,name='secondary-shs');v.r(16.6,-18.4,16.8,16.8,None,GREY)
v.r(225,-13,14,10,None,GOLD,True);v.l(225,-14,239,-14,GOLD,1.2)
v.l(232,-17,225,-36,GOLD,.7,True);v.l(232,-17,239,-36,GOLD,.7,True)
v.dim(0,3,250,3,-12);v.dim(250,-22,250,3,13)
cloud(s,49,183,222, 40,'TBC #16 — LED housing drawn as an indicative envelope. Confirm housing, fixing and setback against the selected product. No drilling schedule released.')
txt(s,64,88,210,15,'01  TRAY + DOWNSTAND        02  SUPPORT SHS        03  REMOVABLE LED CHANNEL',8,GREY)
leader(s,76,166,28,127,'02 / 20 SHS support')
leader(s,188,164,209,139,'01 / 25 downstand',65)
leader(s,180,163,207,173,'03 / diffuser + cable',68)

s=sheet('Mesh / lighting / service access','AP-14 · removable components and segregated service routes','1:5 / NTS',[
 ('MESH PANEL','2200–2600 front infill only. Approx. 20 × 40 diamond mesh; gauge, strand width and open area TBC #16.'),('EDGE / REMOVAL','Capture cut mesh edges in removable perimeter frame. Screw-fix accessible panel to main structure; no sharp exposed edges.'),('LED DRIVERS','Locate in accessible ventilated cabinet service zone. Keep maintenance reachable without dismantling the display.'),('ELECTRICAL','Separate mains/data/low-voltage routes as required by consultant. Driver rating, cable sizes, ventilation and circuits remain #16.')])
v=View(s,44,153,5);v.r(0,0,664,400,None,INK,name='mesh-zone')
for x in range(0,664,40):v.l(x,0,min(x+200,664),400,GREY,.35)
for x in range(0,664,40):v.l(x,400,min(x+200,664),0,GREY,.35)
for x in (20,644):
    for y in (20,380):a,b=v.p(x,y);circle(s,a,b,1,RED)
v.dim(0,0,664,0,7);v.dim(0,0,0,400,-9)
cloud(s,24,178,254, 70,'TBC #16 — PANEL + SERVICE STRATEGY\n\nWorkshop-welded steel subframe → deburr and coat → demountable mesh cassette.\nAccessible mains isolation → ventilated driver zone → protected LV route in upright → shelf channels.\nFixing centres, panel gauge and ventilation area are unapproved. Diagram is not a product specification.')

s=sheet('Storefront head / structural interface','AP-15 · glazing head and pivot reinforcement strategy','1:2 — indicative assembly',[
 ('DOOR HEIGHT','2580 door within 2600 opening leaves 20 total vertical allowance. Propose 10 top / 10 bottom; glazier to confirm #16.'),('LOAD PATH','Top pivot fixed to engineered steel support tied to verified structure. Plaster / gypsum and fascia are not structural supports.'),('FASCIA ACCESS','Removable internal access panel for pivot inspection and sign driver. Panel size and backing construction TBC #16/#26.'),('GLASS RELEASE','Confirm door mass, hardware capacity, deflection limits, hole positions and approved glass fabrication drawings before cutting.')])
v=View(s,85,174,2);v.r(-50,40,180,70,BEIGE,RED,True);v.r(0,10,80,30,None,RED,True)
v.r(34,-40,12,40,None,BLUE,name='door-glass-head');v.l(40,0,40,38,RED,1.5)
v.dim(34,-20,46,-20,9,'12 glass');v.t(-60,130,'VERIFIED STRUCTURE — TYPE / DEPTH TBC',130,8,RED)
leader(s,105,158,172,155,'Pivot to reinforced head',102)
leader(s,136,137,172,127,'Accessible support / fixings',102)
cloud(s,35,199,242, 40,'TBC #16 / #26 — reinforcement, pivot housing and substrate shown diagrammatically. No anchor size or structural capacity specified. Head detail coordinates with AP-16 / AP-17.')

s=sheet('Glazed entrance / jamb and pivot','AP-16 · horizontal joint chain and operation envelope','1:5',[
 ('WIDTH CHAIN','4 + 488 + 4 + 1000 + 4 = 1500. Nominal gaps are design allowances, not approved glass fabrication dimensions.'),('PIVOT','100 from right door edge, viewed externally. Inward opening. Long-side sweep radius 900; short tail radius 100.'),('SEALS / EDGES','Compatible removable gaskets, setting blocks and edge protection. Seal thickness and compression to glazier details #16.'),('COORDINATION','Confirm effective clear opening, short-tail sweep, handles and stops. 1000 leaf width does not certify a 1000 clear escape opening.')])
v=View(s,24,116,5)
# Entire 1500 opening would be 300 paper mm; split into two adjacent detail fields.
v.r(0,0,488,12,None,BLUE);v.r(492,0,180,12,None,BLUE)
v.dim(0,12,488,12,-10);v.dim(488,0,492,0,10,'4 gap')
v.t(0,-100,'FIXED PANE → DOOR JUNCTION / PARTIAL LEAF',135)
v=View(s,180,204,5);v.r(0,0,200,12,None,BLUE);v.l(100,-45,100,50,RED,1,True);v.dim(100,12,200,12,-9)
v.t(-100,-95,'PIVOT ZONE / RIGHT EDGE',90)
cloud(s,24,231,254, 25,'TBC #16 — supplier to resolve pane edge support, seal pockets, pivot footprint and operating clearances before glass order.')

s=sheet('Flush threshold / floor interface','AP-17 · accessible level transition and recessed hardware','1:2 — indicative assembly',[
 ('DESIGN LEVEL','No step or raised sill. Internal and external finished surfaces meet flush; external levels require survey #26.'),('FLOOR PIVOT','Recessed pivot/closer supported by verified structural substrate. Independent of tile and adhesive; removable cover flush with FFL.'),('MOVEMENT / SEAL','Maintain 5 perimeter movement allowance where applicable. Compatible flexible seals; do not bridge joints with rigid grout.'),('SITE CONFIRMATION','Slab depth, services, damp protection and external drainage must be verified before recessing hardware or setting threshold level.')])
v=View(s,65,140,2);v.l(-50,0,350,0,INK,1.2);v.r(-50,-15,400,15,BEIGE,RED,True)
v.r(80,-90,110,75,None,RED,True);v.r(125,10,12,120,None,BLUE)
v.l(131,-50,131,10,RED,1,True);v.t(-45, 20,'FFL ±0 — VERIFY ON SITE #26',130,9,RED)
v.t(65,-110,'RECESSED HARDWARE ENVELOPE TBC',130,8,RED)
leader(s,123,158,184,172,'Independent structural support',93)
leader(s,130,135,184,105,'12 glass / seal TBC #16',93)
cloud(s,27,222,250, 30,'TBC #16 / #26 — tile/adhesive/slab thicknesses and hardware pocket are indicative. Preserve the flush finished level; coordinate seals and maintenance access.')

s=sheet('Halo-lit sign / demountable backing','AP-18 · external logo section and maintenance access','1:2 — indicative assembly',[
 ('BRAND / FINISH','1000 W black logo; uniform native SVG geometry. Halo-lit, 3000 K visual intent. No brass elements.'),('STAND-OFF','25 stand-off behind logo. Concealed threaded spacers to engineered backing; fastener size and count remain #16.'),('POWER / ACCESS','Protected cable penetration to accessible driver on removable internal backing panel. Strain relief and isolation accessible.'),('FABRICATION','Confirm return depth, LED spacing, weather exposure, fixing template, driver enclosure and backing capacity with sign specialist.')])
v=View(s,93,187,2);v.r(0,0,8,220,BEIGE,RED,True);v.r(33,70,12,100,INK)
for y in (90,150):v.l(8,y,33,y,INK,1.5)
v.l(8,120,33,120,GOLD,1,True);v.dim(8,170,33,170,-10,'25 stand-off')
v.t(-100,255,'STRUCTURAL BACKING / ACCESS PANEL TBC',160,9,RED)
leader(s,109,128,153,108,'Threaded spacer / concealed stud',125)
leader(s,102,127,153,146,'Protected wire to accessible driver',125)
leader(s,96,165,153,185,'Removable backing panel',125)
cloud(s,31,219,246, 30,'TBC #16 / #26 — backing thickness and logo return depth are indicative. Show a removable driver-access panel; do not permanently seal the service compartment.')

s=sheet('Rear wall / display-to-wall connection','AP-19 · maintainable back-panel and service junction','1:5 — indicative assembly',[
 ('CENTRAL ZONE','900 design zone, 875 clear between shared 25 uprights. 600 logo leaves 137.5 each side within that clear width.'),('NON-ILLUMINATED','Rear logo remains black and non-illuminated. Cable routes serve adjacent shelf LEDs only.'),('WALL FIXING','Floor/wall-fixed structure. Adjustable brackets to verified substrate; isolate finish from structural fixings.'),('MAINTENANCE','Removable central light-neutral metal panel on accessible clips/screws. Panel gauge, fasteners and reveals TBC #16/#26.')])
v=View(s,57,184,5);v.r(0,0,900,500,BEIGE,RED,True)
for x in (-12.5,887.5):v.r(x,0,25,500,INK)
v.logo(150,180,600);v.dim(0,0,900,0,8)
cloud(s,24,217,253, 35,'TBC #16 / #26 — this is a partial-height panel strategy. Confirm actual panel extent, reveal and wall finish build-up. Separate demountable metal from plaster; no load transfer to gypsum ceiling.')

s=sheet('Counter / plan, elevation and section','AP-20 · 1000 × 450 × 900H retained','1:10')
v=View(s, 30,120,10);v.r(0,0,1000,450,BEIGE,name='counter-top');v.dim(0,0,1000,0,8);v.dim(0,0,0,450,-8);v.t(0,560,'01  PLAN — 1000 × 450',100,10)
v=View(s,165,182,10);v.r(0,0,1000,900,BEIGE,name='counter-front');v.r(350,150,300,450,None,INK,name='mesh-insert');v.dim(0,0,1000,0,8);v.dim(0,0,0,900,-8);v.t(0,1010,'02  FRONT ELEVATION',100,10)
v=View(s,325,182,10);v.r(0,0,450,900,None,INK,name='counter-section');v.r(0,880,450,20,BEIGE,RED,True)
for y in (150,450,700):v.l(20,y,400,y,RED,.6,True)
v.l(400,150,400,880,RED,.6,True);v.dim(0,0,450,0,8);v.t(-50,1010,'03  SECTION',90,10)
cloud(s, 20,218,380, 38,'TBC #16 / #26 — counter interior is a proposed arrangement: 20 top, steel carcass/frame, removable storage trays, rear service chase and adjustable feet. Confirm all new thicknesses, storage layout, service clearances and feet before fabrication. External 1000 × 450 × 900 and 300 × 450 mesh insert remain confirmed.')

s=sheet('Counter / services and removable access','AP-21 · section through proposed service compartment','1:5 — indicative assembly',[
 ('TOP / FRONT','Simple light-coloured top; light-neutral steel front. Proposed 20 solid-surface top on steel support; material/thickness approval #16.'),('STORAGE','Removable steel trays/drawers, avoiding timber or heavy joinery. Slide specification, payload and usable depth TBC #16.'),('POWER / DATA','Proposed rear grommet to segregated service chase; accessible isolation and removable panel on staff side.'),('LEVELLING','Concealed adjustable feet bear on verified floor. Finish to remain at 900H. Foot travel and point loads require consultant confirmation.')])
v=View(s,91,248,5);v.r(0,0,450,900,None,INK);v.r(0,880,450,20,BEIGE,RED,True)
for x in (0,430):v.r(x,100,20,780,INK)
for y in (200,450,700):v.r(20,y,330,20,None,RED,True)
v.l(370,120,370,880,RED,.6,True);v.r(375,180,55,300,None,RED,True)
for x in (30,390):v.r(x,20,25,80,None,RED,True)
v.dim(0,0,450,0,8);v.dim(0,0,0,900,-12)
cloud(s,194,124,87,80,'TBC #16 / #26\n\nInternal dimensions are proposed.\nConfirm selected grommet, service chase, drawer slides and feet.\nStaff access remains a separate open issue.')

s=sheet('Finish junctions / controlled terminations','AP-22 · four typical coordination details','1:2 — indicative assemblies')
for x,title in [(20,'01  FLOOR / WALL'),(120,'02  WALL / CEILING'),(220,'03  DISPLAY / WALL'),(320,'04  INTERNAL CORNER')]:txt(s,x,61,95,10,title,10,bold=True)
v=View(s,30,165,2);v.r(0,0,18,130,BEIGE,RED,True);v.r(23,0,130,10,BEIGE,RED,True);v.dim(18,0,23,0,8,'5 joint');v.t(-20,-90,'Flexible compatible seal\nBuild-up TBC',93,9,RED)
v=View(s,130,165,2);v.r(0,0,18,130,BEIGE,RED,True);v.r(18,112,130,18,BEIGE,RED,True);v.l(18,112,55,112,RED,1,True);v.t(-20,-90,'Controlled reveal TBC\nIndependent ceiling support',93,9,RED)
v=View(s,230,165,2);v.r(0,0,18,130,BEIGE,RED,True);v.r(45,10,25,100,INK);v.l(18,60,45,60,RED,1.2,True);v.t(-20,-90,'Adjustable wall bracket\nAccessible fastener / packer',93,9,RED)
v=View(s,330,165,2);v.r(0,0,18,130,BEIGE,RED,True);v.r(18,0,120,18,BEIGE,RED,True);v.l(18,18,27,27,RED,.7,True);v.t(-20,-90,'Plan / internal return\nReinforced plaster corner\nMovement treatment TBC',93,9,RED)
cloud(s, 20,226,380, 30,'TBC #16 / #26 — all build-ups and bracket/reveal geometries shown provisionally. Retain the confirmed 5 floor perimeter allowance. Confirm movement, substrate, finish stops, fixing capacity and visible sample before construction.')

s=sheet('Technical specification / coordinated schedule','AP-23 · performance decisions remain with the responsible specialists','NTS')
rows=[('M01','Main / secondary steel','25×25×1.6 / 20×20×1.6 SHS; matte black PPC','Weld/anchor design, corrosion prep and loading #16'),('M02','Doors / neutral panels','1.5 mm PPC steel doors; flush, handleless','Hardware, reveals and durability sample #16'),('M03','Wall / ceiling','Raw light plaster; smooth off-white gypsum ceiling','Substrate, thickness and fire requirements #26/#16'),('M04','Floor / sealants','1200×600 porcelain; 2 grout; 5 perimeter','Adhesive, slip suitability, sealant compatibility #16/#26'),('G01','Glazing / hardware','1000×2580×12 toughened pivot leaf; 488 fixed pane','Fixed pane thickness, gaps, hardware capacity #16'),('M05','Mesh','Approx. 20×40 diamond; black; framed removable edges','Gauge, perimeter channel and screw centres #16'),('M06','Shelf','3 mm MS tray; 25 downstand; 250 deep','Load, tray fit and LED extrusion selection #16'),('L01','Lighting / access','3000 K; two tracks; shelf and plinth LEDs','Photometry, wattage, drivers, circuits, ventilation #16'),('S01','Logo / fixings','External halo-lit at 25 stand-off; internal unlit','Backing, studs, maintenance opening and IP needs #16'),('C01','Counter / joints','1000×450×900; 300×450 mesh; light top','Proposed 20 top, storage, feet, grommet, reveals #16')]
for x,w,t in [(20,18,'CODE'),( 40,75,'ASSEMBLY'),(115,145,'DESIGN BASIS'),(265,135,'CONFIRM BEFORE FABRICATION')]:txt(s,x,58,w,8,t,8,GOLD,bold=True)
for i,(a,b,c,d) in enumerate(rows):
    y=73+i*18;line(s,20,y-2,400,y-2,BEIGE,.5)
    for x,w,t in [(20,18,a),(40, 70,b),(115,145,c),(265,135,d)]:txt(s,x,y,w,16,t,9)

s=sheet('Installation / quality and maintenance sequence','AP-24 · hold points before release to fabrication','NTS')
steps=[('01 / SURVEY','Confirm finished envelope, FFL, substrate, wall construction and site north. Record measured values separately from design dimensions.'),('02 / CONSULTANT','Resolve 450 side routes; assess pivot clear opening, fire / MEP coordination, structure and fixing loads. Sign off TBC #16.'),('03 / SHOP DRAWINGS','Fabricator resolves cut sizes, shared posts, welds, glass gaps, tolerances and adjustment. Release supplier templates only after approval.'),('04 / SAMPLE + FABRICATE','Approve finishes, edge conditions and light concealment. Workshop weld and coat; protect completed finishes during transport.'),('05 / SITE ASSEMBLY','Set out → level → anchor to verified structure → route/test services → install trays, panels, glass and logo → adjust and protect.'),('06 / COMMISSION','Test door sweep, access, lighting and isolation. Demonstrate panel removal; hand over wiring, maintenance instructions and surveyed as-builts.')]
for i,(a,b) in enumerate(steps):
    x=22+(i%2)*195;y= 60+(i//2)*65
    txt(s,x,y,179,9,a,12,GOLD,bold=True);txt(s,x,y+14,178,40,b,11)

s=sheet('Open decisions / consultant coordination','AP-25 · retain clouds until documented approval','NTS')
decisions=[('#4','Finished envelope','Survey 2400 × 3106 after plaster; check diagonals, plumb and ceiling height.','Surveyor / owner'),('#16','Counter / escape','450 side gaps vs 900 project assumption. Confirm compliant staff/customer route; retain centred design pending decision.','Fire / access consultant'),('#16','Glass / hardware','2580 door height leaves 20 vertical allowance. Confirm split, fixed-pane thickness, operating gaps and pivot capacity.','Glazier / engineer'),('#16','Services / fixings','No anchor capacities, lighting circuits or driver ventilation established. Review proposed AP-13 to AP-22 assemblies.','MEP / structural / fabricator'),('#23','Orientation','North toward storefront is assumed. No surveyed north is represented in this package.','Surveyor'),('#26','Wall / datum / finishes','100 wall is historical assumption; plans show finish lines only. Verify wall/ceiling/floor build-ups, levels and movement joints.','Surveyor / architect')]
for i,(a,b,c,d) in enumerate(decisions):
    y=60+i*32;cloud(s,20,y,380,26,f'{a}  {b.upper()}  /  {d}\n{c}')

s=sheet('Drawing register / dimensional close-out','AP-26 · presentation consistency audit and next release','NTS')
txt(s,20,58,185,10,'SHEET REGISTER',11,GOLD,bold=True)
for i,(code,title,scale) in enumerate(index):
    col=i//13;row=i%13;txt(s,20+col*196,74+row*12,185,11,f'{code}  {title}',8.5)
txt(s,20,237,380,23,'CHECKED: 450 + 1500 + 450 = 2400  |  488 + 1000 + 3×4 = 1500  |  450 + 4×664 = 3106\n100 + 500 + 5×400 = 2600  |  3106 − 250 − 2156 = 700  |  counter 1000×450×900\nSurveys and fabrication decisions remain open. Detailed read-back audit accompanies this presentation.',9,GREY)

prs.core_properties.title='BORN FRAGRANCE | Architectural Interior Presentation'
prs.core_properties.subject='Client and consultant coordination; proposed technical details'
prs.core_properties.author='ChatGPT / Codex'
for slide in prs.slides:
    for shape in slide.shapes:
        # Explicitly suppress the default theme's shadow on drafting primitives.
        style=shape._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
        if style is not None:
            shape._element.remove(style)
        sppr=shape._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}spPr')
        if sppr is not None:
            sppr.append(OxmlElement('a:effectLst'))
prs.save(OUT/'BORN_FRAGRANCE_Architecture.pptx')
(OUT/'geometry_manifest.json').write_text(json.dumps({'source_commit':SHA,'spec_sha256':hashlib.sha256((ROOT/'spec/SHOP_SPEC.json').read_bytes()).hexdigest(),'checks':checks,'sheets':index},indent=2))
print(f'Built {len(prs.slides)} A3 slides with editable vector geometry')
